"""Mentor comments, the resident's reply, the mentor's resolve and approval, and the report order."""
import io
import os

os.environ.setdefault("FERNET_KEY", "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA=")

from docx import Document
from fastapi.testclient import TestClient

from api.main import app
from tests.test_share_lifecycle import _create_project, _docx_text, _register, _seed_analysis_package


def _share(client, project_id):
    response = client.post(f"/share/{project_id}/create", json={"mentor_email": "mentor@example.com"})
    assert response.status_code == 200, response.text
    return response.json()["token"]


def _mentor():
    return TestClient(app)


def test_comment_reply_resolve_and_approval_cycle_ends_the_report(client):
    _register(client, "cycle-owner@example.com")
    project = _create_project(client)
    _seed_analysis_package(project["id"])
    token = _share(client, project["id"])

    mentor = _mentor()
    created = mentor.post(f"/share/view/{token}/comment", json={"author_name": "Dr Rivera", "text": "Please state how missing values were handled."})
    assert created.status_code == 200, created.text
    comment_id, edit_token = created.json()["id"], created.json()["edit_token"]

    review = client.get(f"/share/{project['id']}/review").json()
    assert [c["text"] for c in review["comments"]] == ["Please state how missing values were handled."]
    assert review["comments"][0]["addressed_at"] is None and review["approval"] is None

    reply = client.post(f"/share/{project['id']}/comment/{comment_id}/reply", json={"reply": "Added a sentence to the methods."})
    assert reply.status_code == 200, reply.text
    seen_by_mentor = mentor.get(f"/share/view/{token}").json()
    assert seen_by_mentor["comments"][0]["resident_reply"] == "Added a sentence to the methods."
    assert seen_by_mentor["comments"][0]["addressed_at"] and not seen_by_mentor["comments"][0]["resolved_at"]

    assert mentor.patch(f"/share/view/{token}/comment/{comment_id}/resolve", json={"edit_token": "wrong", "resolved": True}).status_code == 403
    resolved = mentor.patch(f"/share/view/{token}/comment/{comment_id}/resolve", json={"edit_token": edit_token, "resolved": True})
    assert resolved.status_code == 200 and resolved.json()["resolved_at"]

    approval = mentor.post(f"/share/view/{token}/approve", json={"author_name": "Dr Rivera", "note": "Good work."})
    assert approval.status_code == 200, approval.text
    assert mentor.get(f"/share/view/{token}").json()["approval"]["approved_by_name"] == "Dr Rivera"
    assert client.get(f"/share/{project['id']}/review").json()["approval"]["approved_note"] == "Good work."

    text = _docx_text(client.get(f"/report/project/{project['id']}/docx").content)
    assert text.index("Audit Trail") < text.index("Mentor Feedback")
    tail = text[text.index("Mentor Feedback"):]
    assert "Resident reply" in tail and "Added a sentence to the methods." in tail
    assert "Status: Resolved by the mentor" in tail
    assert "Mentor approval: approved by Dr Rivera" in tail and "Good work." in tail

    doc = Document(io.BytesIO(client.get(f"/report/project/{project['id']}/docx").content))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    assert paragraphs[-1].startswith("Mentor approval:")  # the mentor has the last word
    first = next(i for i, p in enumerate(doc.paragraphs) if p.text == "Mentor Feedback")
    assert all(p.alignment is None for p in doc.paragraphs[first:])  # only the report title is centered


def test_a_new_reply_reopens_a_resolved_comment_and_only_the_owner_can_reply(client):
    _register(client, "cycle-owner2@example.com")
    project = _create_project(client)
    token = _share(client, project["id"])
    mentor = _mentor()
    created = mentor.post(f"/share/view/{token}/comment", json={"author_name": "Dr B", "text": "Check the denominator."}).json()
    mentor.patch(f"/share/view/{token}/comment/{created['id']}/resolve", json={"edit_token": created["edit_token"], "resolved": True})
    client.post(f"/share/{project['id']}/comment/{created['id']}/reply", json={"reply": "Fixed again."})
    assert mentor.get(f"/share/view/{token}").json()["comments"][0]["resolved_at"] is None

    stranger = _mentor()
    assert stranger.post(f"/share/{project['id']}/comment/{created['id']}/reply", json={"reply": "x"}).status_code in (401, 403)
    assert stranger.get(f"/share/{project['id']}/review").status_code in (401, 403)


def test_approval_note_loses_control_characters_and_needs_a_name(client):
    _register(client, "cycle-owner3@example.com")
    project = _create_project(client)
    token = _share(client, project["id"])
    mentor = _mentor()
    assert mentor.post(f"/share/view/{token}/approve", json={"author_name": "", "note": "x"}).status_code == 422
    ok = mentor.post(f"/share/view/{token}/approve", json={"author_name": "Dr\x00 C", "note": "Nice\x01 job"}).json()
    assert ok["approved_by_name"] == "Dr C" and ok["approved_note"] == "Nice job"


def test_pdf_report_also_ends_with_the_mentor_review(client):
    from tests.test_report import _pdf_text

    _register(client, "cycle-owner4@example.com")
    project = _create_project(client)
    _seed_analysis_package(project["id"])
    token = _share(client, project["id"])
    mentor = _mentor()
    mentor.post(f"/share/view/{token}/comment", json={"author_name": "Dr D", "text": "Looks right."})
    mentor.post(f"/share/view/{token}/approve", json={"author_name": "Dr D", "note": "Good work."})
    text = _pdf_text(client.get(f"/report/project/{project['id']}/pdf").content)
    assert text.index("Audit Trail") < text.index("Mentor Feedback") < text.index("Mentor approval")
