"""Focused share, mentor comment, and notification lifecycle tests."""
import io
import json
import os
from datetime import date, datetime, timedelta

os.environ.setdefault("FERNET_KEY", "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA=")

from docx import Document
from fastapi.testclient import TestClient

from api.database import SessionLocal
from api.main import app
from api.models_db import AnalysisRun, AuditLog, EditHistory, MentorComment, MentorShare, NotificationDelivery, Project, Upload
from api.services import notifications as notification_service

PASSWORD = "password123"
TINY_PNG_BASE64 = (
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwC"
    "AAAAC0lEQVR42mP8/x8AAwMCAO+/p9sAAAAASUVORK5CYII="
)


def _register(test_client, email):
    response = test_client.post("/auth/register", json={"email": email, "password": PASSWORD})
    assert response.status_code == 200, response.text
    return response.json()


def _create_project(test_client, title="Mentor Review Project", deadline=None):
    response = test_client.post(
        "/projects",
        json={"title": title, "description": "Improve discharge reliability", "deadline": deadline},
    )
    assert response.status_code == 200, response.text
    return response.json()


def _docx_text(content: bytes) -> str:
    doc = Document(io.BytesIO(content))
    parts = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def _seed_analysis_package(project_id, *, title="Mentor-facing final title"):
    db = SessionLocal()
    now = datetime.utcnow()
    upload = Upload(
        project_id=project_id,
        filename="analysis.csv",
        original_filename="mentor-analysis-source.csv",
        encrypted_path="/tmp/mentor-analysis.enc",
        quality_flags=json.dumps([
            {"col": "wait_days", "rule": "check_missing", "severity": "WARNING", "msg": "wait_days is 12% missing"}
        ]),
        acknowledged_flags=json.dumps([
            {"col": "wait_days", "rule": "check_missing", "severity": "WARNING", "msg": "wait_days is 12% missing"}
        ]),
    )
    db.add(upload)
    db.flush()
    run = AnalysisRun(
        project_id=project_id,
        upload_id=upload.id,
        template="run_chart",
        parameters=json.dumps({"date_col": "month", "value_col": "wait_days"}),
        result_json=json.dumps(
            {
                "methods": "Run chart methods for mentor review.",
                "result_summary": "Median wait time fell from 10 to 7 days.",
                "interpretation": "AI interpretation before resident review.",
                "figure_base64": TINY_PNG_BASE64,
                "table": [{"period": "Baseline", "median": 10}, {"period": "Follow-up", "median": 7}],
            }
        ),
        code_r="# mentor report R code",
        code_spss="* mentor SPSS code",
        code_sas="/* mentor SAS code */",
        created_at=now,
    )
    db.add(run)
    db.add_all(
        [
            EditHistory(
                project_id=project_id,
                field="title",
                original_text="Mentor Review Project",
                edited_text=title,
                timestamp=now - timedelta(minutes=3),
            ),
            EditHistory(
                project_id=project_id,
                field="caption",
                original_text="",
                edited_text="Mentor-facing figure caption",
                timestamp=now - timedelta(minutes=2),
            ),
            EditHistory(
                project_id=project_id,
                field="interpretation",
                original_text="AI interpretation before resident review.",
                edited_text="Resident interpretation prepared for mentor review.",
                timestamp=now - timedelta(minutes=1),
            ),
        ]
    )
    db.commit()
    run_id = run.id
    db.close()
    return run_id


def test_share_create_accepts_json_expiration_sends_invite_once_and_reuses_active_mentor_share(client, monkeypatch):
    _register(client, "admin@example.com")
    project = _create_project(client, deadline="2030-03-15")
    sent_invites = []

    def fake_invite(email, project_title, share_url, deadline=None):
        sent_invites.append({"email": email, "project_title": project_title, "share_url": share_url, "deadline": deadline})

    monkeypatch.setattr(notification_service, "send_share_invite", fake_invite)
    expires_at = (datetime.utcnow() + timedelta(days=5)).replace(microsecond=0)

    created = client.post(
        f"/share/{project['id']}/create",
        json={"mentor_email": "Mentor@Example.com", "expires_at": expires_at.isoformat()},
    )
    assert created.status_code == 200, created.text
    body = created.json()
    assert body["mentor_email"] == "mentor@example.com"
    assert body["expires_at"].startswith(expires_at.isoformat())
    assert body["notification_status"] == "sent"
    assert len(sent_invites) == 1
    assert sent_invites[0]["email"] == "mentor@example.com"
    assert sent_invites[0]["project_title"] == project["title"]
    assert body["token"] in sent_invites[0]["share_url"]
    assert sent_invites[0]["deadline"] == "2030-03-15"

    repeated = client.post(f"/share/{project['id']}/create", json={"mentor_email": "mentor@example.com"})
    assert repeated.status_code == 200, repeated.text
    assert repeated.json()["token"] == body["token"]
    assert repeated.json()["notification_status"] == "existing_share"
    assert len(sent_invites) == 1

    with SessionLocal() as db:
        share = db.query(MentorShare).filter_by(token=body["token"]).one()
        deliveries = db.query(NotificationDelivery).filter_by(share_id=share.id, kind="share_invite").all()
        assert len(deliveries) == 1
        assert deliveries[0].recipient_email == "mentor@example.com"
        assert deliveries[0].status == "sent"
        assert deliveries[0].sent_at is not None


def test_share_create_records_failed_notification_without_blocking_review_link(client, monkeypatch):
    _register(client, "admin@example.com")
    project = _create_project(client)

    def failing_invite(email, project_title, share_url, deadline=None):
        raise RuntimeError("SMTP unavailable")

    monkeypatch.setattr(notification_service, "send_share_invite", failing_invite)

    response = client.post(f"/share/{project['id']}/create", json={"mentor_email": "mentor@example.com"})
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["token"]
    assert body["notification_status"] == "failed"

    with SessionLocal() as db:
        share = db.query(MentorShare).filter_by(token=body["token"]).one()
        delivery = db.query(NotificationDelivery).filter_by(share_id=share.id, kind="share_invite").one()
        assert delivery.status == "failed"
        assert delivery.recipient_email == "mentor@example.com"
        assert delivery.error_message == "Notification delivery failed"
        assert delivery.sent_at is None
        assert db.query(AuditLog).filter_by(project_id=project["id"], action="notification_failed").count() == 1


def test_share_revoke_regenerate_and_expired_or_revoked_tokens_are_rejected(client, monkeypatch):
    _register(client, "admin@example.com")
    project = _create_project(client)
    monkeypatch.setattr(notification_service, "send_share_invite", lambda *args, **kwargs: None)

    expired = client.post(
        f"/share/{project['id']}/create",
        json={"expires_at": (datetime.utcnow() - timedelta(days=1)).isoformat()},
    )
    assert expired.status_code == 200, expired.text
    assert client.get(f"/share/view/{expired.json()['token']}").status_code == 404

    active = client.post(f"/share/{project['id']}/create", json={"mentor_email": "revoke@example.com"})
    assert active.status_code == 200, active.text
    revoked = client.post(f"/share/{project['id']}/revoke", json={"token": active.json()["token"]})
    assert revoked.status_code == 200, revoked.text
    assert revoked.json()["revoked_at"] is not None
    assert client.get(f"/share/view/{active.json()['token']}").status_code == 404

    to_regenerate = client.post(f"/share/{project['id']}/create", json={"mentor_email": "regen@example.com"})
    assert to_regenerate.status_code == 200, to_regenerate.text
    replacement = client.post(f"/share/{project['id']}/regenerate", json={"token": to_regenerate.json()["token"]})
    assert replacement.status_code == 200, replacement.text
    assert replacement.json()["token"] != to_regenerate.json()["token"]
    assert client.get(f"/share/view/{to_regenerate.json()['token']}").status_code == 404
    assert client.get(f"/share/view/{replacement.json()['token']}").status_code == 200

    with SessionLocal() as db:
        actions = [row.action for row in db.query(AuditLog).filter_by(project_id=project["id"]).all()]
        assert "share_revoked" in actions
        assert "share_regenerated" in actions

def test_share_revoke_and_regenerate_require_explicit_target_and_do_not_mutate(client, monkeypatch):
    _register(client, "admin@example.com")
    project = _create_project(client)
    monkeypatch.setattr(notification_service, "send_share_invite", lambda *args, **kwargs: None)

    active = client.post(f"/share/{project['id']}/create", json={"mentor_email": "target@example.com"})
    assert active.status_code == 200, active.text
    token = active.json()["token"]

    assert client.post(f"/share/{project['id']}/revoke", json={}).status_code in (400, 422)
    assert client.post(f"/share/{project['id']}/regenerate", json={}).status_code in (400, 422)

    with SessionLocal() as db:
        share = db.query(MentorShare).filter_by(token=token).one()
        assert share.revoked_at is None
        assert db.query(MentorShare).filter_by(project_id=project["id"]).count() == 1

    assert client.get(f"/share/view/{token}").status_code == 200


def test_mentor_view_comment_author_lifecycle_and_scoped_report_download(client, monkeypatch):
    _register(client, "admin@example.com")
    project = _create_project(client, title="Scoped Mentor Package")
    run_id = _seed_analysis_package(project["id"])
    other_project = _create_project(client, title="Other Mentor Package")
    _seed_analysis_package(other_project["id"], title="Other final title")
    monkeypatch.setattr(notification_service, "send_share_invite", lambda *args, **kwargs: None)

    share = client.post(f"/share/{project['id']}/create", json={"mentor_email": "mentor-a@example.com"}).json()
    other_share = client.post(f"/share/{other_project['id']}/create", json={"mentor_email": "mentor-b@example.com"}).json()
    other_comment = client.post(
        f"/share/view/{other_share['token']}/comment",
        json={"author_name": "Other Mentor", "author_email": "other@example.com", "text": "Other project feedback"},
    )
    assert other_comment.status_code == 200, other_comment.text

    view = client.get(f"/share/view/{share['token']}")
    assert view.status_code == 200, view.text
    package = view.json()
    assert package["project"]["title"] == "Scoped Mentor Package"
    assert package["template"] == "run_chart"
    assert package["methods"] == "Run chart methods for mentor review."
    assert package["result_summary"] == "Median wait time fell from 10 to 7 days."
    assert package["interpretation"] == "Resident interpretation prepared for mentor review."
    assert package["table"] == [{"period": "Baseline", "median": 10}, {"period": "Follow-up", "median": 7}]
    assert package["figure_base64"] == TINY_PNG_BASE64
    assert package["caption"] == "Mentor-facing figure caption"
    assert package["limitations"] == [
        {"col": "wait_days", "rule": "check_missing", "severity": "WARNING", "msg": "wait_days is 12% missing"}
    ]
    assert package["code_r"] == "# mentor report R code"
    assert package["code_spss"] == "* mentor SPSS code"
    assert package["code_sas"] == "/* mentor SAS code */"
    assert "report_urls" not in package
    assert client.get(f"/api/share/view/{share['token']}/report/docx").status_code == 404
    assert package["comments"] == []

    added = client.post(
        f"/share/view/{share['token']}/comment",
        json={"author_name": "Dr Mentor", "author_email": "Mentor@Example.com", "text": "Initial scoped feedback"},
    )
    assert added.status_code == 200, added.text
    comment = added.json()
    assert comment["author_email"] == "mentor@example.com"
    assert comment["text"] == "Initial scoped feedback"

    wrong_author = client.patch(
        f"/share/view/{share['token']}/comment/{comment['id']}",
        json={"author_email": "someone-else@example.com", "text": "Hijacked feedback"},
    )
    assert wrong_author.status_code == 403

    edited = client.patch(
        f"/share/view/{share['token']}/comment/{comment['id']}",
        json={"author_email": "MENTOR@example.com", "text": "Scoped edited feedback"},
    )
    assert edited.status_code == 200, edited.text
    assert edited.json()["text"] == "Scoped edited feedback"
    assert edited.json()["updated_at"] is not None

    report = client.get(f"/report/{run_id}/docx")
    assert report.status_code == 200, report.text
    report_text = _docx_text(report.content)
    assert "Scoped edited feedback" in report_text
    assert "Other project feedback" not in report_text

    deleted = client.request(
        "DELETE",
        f"/share/view/{share['token']}/comment/{comment['id']}",
        json={"author_email": "mentor@example.com"},
    )
    assert deleted.status_code == 200, deleted.text
    assert deleted.json() == {"ok": True}
    assert client.get(f"/share/view/{share['token']}").json()["comments"] == []
    with SessionLocal() as db:
        actions = [row.action for row in db.query(AuditLog).filter_by(project_id=project["id"]).all()]
        assert "mentor_comment_created" in actions
        assert "mentor_comment_edited" in actions
        assert "mentor_comment_deleted" in actions
        assert "report_downloaded" in actions


def test_owner_can_edit_and_admin_can_delete_protected_mentor_comments_while_other_residents_are_denied(monkeypatch):
    admin_client = TestClient(app)
    owner_client = TestClient(app)
    other_client = TestClient(app)
    public_client = TestClient(app)
    try:
        _register(admin_client, "admin@example.com")
        _register(owner_client, "owner@example.com")
        _register(other_client, "other@example.com")
        monkeypatch.setattr(notification_service, "send_share_invite", lambda *args, **kwargs: None)
        project = _create_project(owner_client, title="Owner Protected Comments")
        share = owner_client.post(f"/share/{project['id']}/create", json={"mentor_email": "mentor@example.com"}).json()
        added = public_client.post(
            f"/share/view/{share['token']}/comment",
            json={"author_name": "Mentor", "author_email": "mentor@example.com", "text": "Needs owner review"},
        )
        assert added.status_code == 200, added.text
        comment_id = added.json()["id"]

        denied = other_client.patch(
            f"/share/{project['id']}/comment/{comment_id}",
            json={"text": "Other resident should not edit"},
        )
        assert denied.status_code == 403

        owner_edit = owner_client.patch(
            f"/share/{project['id']}/comment/{comment_id}",
            json={"text": "Owner-visible resolution"},
        )
        assert owner_edit.status_code == 200, owner_edit.text
        assert owner_edit.json()["text"] == "Owner-visible resolution"

        admin_delete = admin_client.delete(f"/share/{project['id']}/comment/{comment_id}")
        assert admin_delete.status_code == 200, admin_delete.text
        assert admin_delete.json() == {"ok": True}
        assert public_client.get(f"/share/view/{share['token']}").json()["comments"] == []

        with SessionLocal() as db:
            comment = db.get(MentorComment, comment_id)
            assert comment.deleted_at is not None
    finally:
        admin_client.close()
        owner_client.close()
        other_client.close()
        public_client.close()


def test_admin_deadline_reminder_sends_once_per_active_share_deadline_kind_and_records_deliveries(client, monkeypatch):
    _register(client, "admin@example.com")
    resident_client = TestClient(app)
    try:
        _register(resident_client, "resident@example.com")
        deadline = (date.today() + timedelta(days=3)).isoformat()
        project = _create_project(client, title="Deadline Reminder Project", deadline=deadline)
        now = datetime.utcnow()
        with SessionLocal() as db:
            db.add_all(
                [
                    MentorShare(
                        project_id=project["id"],
                        token="active-deadline-token-a",
                        mentor_email="mentor-a@example.com",
                        created_at=now,
                        expires_at=now + timedelta(days=30),
                    ),
                    MentorShare(
                        project_id=project["id"],
                        token="active-deadline-token-b",
                        mentor_email="mentor-b@example.com",
                        created_at=now,
                        expires_at=now + timedelta(days=30),
                    ),
                    MentorShare(
                        project_id=project["id"],
                        token="revoked-deadline-token",
                        mentor_email="revoked@example.com",
                        created_at=now,
                        expires_at=now + timedelta(days=30),
                        revoked_at=now,
                    ),
                    MentorShare(
                        project_id=project["id"],
                        token="expired-deadline-token",
                        mentor_email="expired@example.com",
                        created_at=now - timedelta(days=40),
                        expires_at=now - timedelta(days=1),
                    ),
                ]
            )
            db.commit()

        sent_reminders = []

        def fake_deadline_reminder(email, reminder_deadline, project_title, share_url):
            sent_reminders.append(
                {"email": email, "deadline": reminder_deadline, "project_title": project_title, "share_url": share_url}
            )

        monkeypatch.setattr(notification_service, "send_deadline_reminder", fake_deadline_reminder)

        denied = resident_client.post("/notifications/deadline-reminders/run")
        assert denied.status_code == 403

        first_run = client.post("/notifications/deadline-reminders/run")
        assert first_run.status_code == 200, first_run.text
        assert first_run.json()["sent"] == 2
        assert first_run.json()["failed"] == 0
        assert {item["email"] for item in sent_reminders} == {"mentor-a@example.com", "mentor-b@example.com"}
        assert {item["deadline"] for item in sent_reminders} == {deadline}
        assert all(item["project_title"] == "Deadline Reminder Project" for item in sent_reminders)
        assert all(item["share_url"].endswith(("/mentor/active-deadline-token-a", "/mentor/active-deadline-token-b")) for item in sent_reminders)

        kind = f"deadline_reminder:{deadline}"
        with SessionLocal() as db:
            deliveries = db.query(NotificationDelivery).filter_by(project_id=project["id"], kind=kind).all()
            assert sorted(delivery.recipient_email for delivery in deliveries) == [
                "mentor-a@example.com",
                "mentor-b@example.com",
            ]
            assert {delivery.status for delivery in deliveries} == {"sent"}
            assert all(delivery.sent_at is not None for delivery in deliveries)
            assert db.query(AuditLog).filter_by(project_id=project["id"], action="notification_sent").count() == 2

        second_run = client.post("/notifications/deadline-reminders/run")
        assert second_run.status_code == 200, second_run.text
        assert second_run.json()["sent"] == 0
        assert second_run.json()["failed"] == 0
        assert len(sent_reminders) == 2
        with SessionLocal() as db:
            assert db.query(NotificationDelivery).filter_by(project_id=project["id"], kind=kind).count() == 2
    finally:
        resident_client.close()
